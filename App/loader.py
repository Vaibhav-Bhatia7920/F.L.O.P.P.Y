from pypdf import PdfReader
from docx import Document
from pathlib import Path
from langchain_text_splitters import RecursiveCharacterTextSplitter  

def classify_files(directory):
    pdf, docx, txt = [], [], []
    for file in directory.iterdir():
        if file.is_file():
            if str(file).endswith('.pdf'):
                pdf.append(file)
            elif str(file).endswith('.docx'):
                docx.append(file)
            elif str(file).endswith('.txt'):
                txt.append(file)
    return pdf, docx, txt

def read_files(pdf_files, docx_files, txt_files):
    pdf_content, docx_content, txt_content = [], [], []
    for pdf in pdf_files:
        reader = PdfReader(pdf)
        text = ''
        for page in reader.pages:
            text += page.extract_text()
        pdf_content.append(text)
    for docx in docx_files:
        doc = Document(docx)
        text = ''
        for para in doc.paragraphs:
            text += para.text + '\n'
        docx_content.append(text)
    for txt in txt_files:
        with open(txt, 'r') as f:
            txt_content.append(f.read())
    return pdf_content, docx_content, txt_content


def chunk_files(directory):
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
    all_chunks = []
    for file in directory.iterdir():
        if not file.is_file():
            continue
        
        text = ""
        print(file)
        if file.suffix == ".pdf":
            reader = PdfReader(file)
            for page in reader.pages:
                txt += page.extract_text()
                
        elif file.suffix == ".docx":
            doc = Document(file)
            for para in doc.paragraphs:
                text += para.text + "\n"
            file_type = "docx"

        elif file.suffix == ".txt":
            with open(file, "r", encoding="utf-8") as f:
                text = f.read()
            file_type = "txt"

        else:
            continue

        chunks = text_splitter.split_text(text)

        for i, chunk in enumerate(chunks):
            all_chunks.append({
                "text": chunk,
                "source": file.name,
                "type": file_type,
                "path": str(Path(file).resolve()),
                "chunk_id": i,
                "total_chunks": len(chunks)
            })
            print("Hi")
            print(str(Path(file).resolve()))

    return all_chunks

if __name__ == "__main__":
    directory = Path('path/to/your/directory')
    pdf_files, docx_files, txt_files = classify_files(directory)
    pdf_content, docx_content, txt_content = read_files(pdf_files, docx_files, txt_files)       